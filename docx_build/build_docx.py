# -*- coding: utf-8 -*-
"""Build the updated ECShop report DOCX from the MD content using the old DOCX as style template."""
import copy, re, struct, os
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

SRC = r'F:/PKA/Code/DALN/docx_build/BaoCao_DALN.docx'
OUT = r'C:/Users/phapd/Downloads/BaoCao_DALN_HoanChinh.docx'
MD  = open(r'C:/Users/phapd/Downloads/BaoCao_DALN_CapNhat_HoanChinh.md', encoding='utf-8').read()
DIAG = r'F:/PKA/Code/DALN/docx_build/diag'

doc = docx.Document(SRC)
body = doc.element.body

# ---------- capture template pieces ----------
children = list(body.iterchildren())
def para_text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))

cover_els = children[0:7]  # P001..P006 (đến dòng GV), bỏ đoạn '.' cuối để bìa không tràn trang
phieu_start = phieu_end = None
cham_table_idx = None
for i, el in enumerate(children):
    t = para_text(el)
    if phieu_start is None and 'KỲ THI KẾT THÚC HỌC PHẦN' in t:
        phieu_start = i
    if el.tag == qn('w:tbl') and 'Giảng viên chấm thi 1' in para_text(el):
        cham_table_idx = i
for i, el in enumerate(children):
    if el.tag == qn('w:p') and el.find(qn('w:pPr')) is not None and el.find(qn('w:pPr')).find(qn('w:sectPr')) is not None:
        sect_break_el = copy.deepcopy(el)
        break

# phiếu block: from phieu_start through the 1x2 table + 1 trailing paragraph
phieu_end = cham_table_idx + 1
phieu_els = children[phieu_start:phieu_end + 1]

sect1 = doc.sections[0].sectPr if hasattr(doc.sections[0], 'sectPr') else None
# collect both sectPr for header fix
sectprs = []
for s in doc.sections:
    sectprs.append(s)

# ---------- clear body ----------
last_sectpr = body.find(qn('w:sectPr'))
for el in list(body.iterchildren()):
    if el is not last_sectpr:
        body.remove(el)

def append_el(el):
    body.insert(list(body).index(last_sectpr) if last_sectpr is not None else len(body), el)
    # simpler: insert before last_sectpr
def add_el(el):
    if last_sectpr is not None:
        last_sectpr.addprevious(el)
    else:
        body.append(el)

# ---------- helpers ----------
def set_run(r, size=14, bold=False, italic=False, mono=False, sub=False, color=None):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if mono:
        r.font.name = 'Consolas'
        r._element.rPr.rFonts.set(qn('w:cs'), 'Consolas')
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    if sub:
        r.font.subscript = True
    if color:
        r.font.color.rgb = RGBColor.from_string(color)

def para(text='', size=14, bold=False, italic=False, align=None, indent=Cm(0.75),
         before=0, after=8, line=1.5, mono=False, color=None, page_break=False,
         keep_next=False, hanging=False):
    p = doc.add_paragraph()
    add_el(p._p)
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if indent is not None:
        if hanging:
            pf.left_indent = indent
            pf.first_line_indent = -indent
        else:
            pf.first_line_indent = indent
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if page_break:
        pf.page_break_before = True
    if keep_next:
        pf.keep_with_next = True
    if text:
        add_inline(p, text, size=size, bold=bold, italic=italic, mono=mono, color=color)
    return p

TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|\$[A-Za-z]_\{[^}]+}\$)')

def add_inline(p, text, size=14, bold=False, italic=False, mono=False, color=None):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            inner = tok[2:-2]
            for sub in re.split(r'(`[^`\n]+`)', inner):
                if not sub:
                    continue
                if sub.startswith('`') and sub.endswith('`'):
                    r = p.add_run(sub[1:-1]); set_run(r, max(size - 2, 10), True, italic, True, color=color)
                else:
                    r = p.add_run(sub); set_run(r, size, True, italic, mono, color=color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run(r, max(size - 2, 10), bold, italic, True, color=color)
        elif tok.startswith('$') and tok.endswith('$'):
            m = re.match(r'\$([A-Za-z])_\{([^}]+)}\$', tok)
            if m:
                r = p.add_run(m.group(1)); set_run(r, size, bold, italic, mono)
                r2 = p.add_run(m.group(2)); set_run(r2, size, bold, italic, mono, sub=True)
            else:
                r = p.add_run(tok.strip('$')); set_run(r, size, bold, italic, mono)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run(r, size, bold, True, mono, color=color)
        else:
            r = p.add_run(tok); set_run(r, size, bold, italic, mono, color=color)
    return p

def png_size(path):
    with open(path, 'rb') as f:
        head = f.read(26)
    w, h = struct.unpack('>II', head[16:24])
    return w, h

def add_figure(name, landscape=False, max_w=15.0, max_h=22.5):
    if landscape:
        max_w, max_h = 23.4, 15.0
    path = os.path.join(DIAG, name + '.png')
    w, h = png_size(path)
    wi = max_w
    hi = wi * h / w
    if hi > max_h:
        hi = max_h
        wi = hi * w / h
    p = doc.add_paragraph()
    add_el(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(wi), height=Cm(hi))

def section_break_para(landscape=False, restart=False):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    sp = copy.deepcopy(last_sectpr)
    if landscape:
        pgsz = sp.find(qn('w:pgSz'))
        pgsz.set(qn('w:w'), '16838')
        pgsz.set(qn('w:h'), '11906')
        pgsz.set(qn('w:orient'), 'landscape')
    # chỉ section bắt đầu phần thân mới reset số trang về 1
    for pgnt in sp.findall(qn('w:pgNumType')):
        sp.remove(pgnt)
    if restart:
        sp.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>'))
    pPr.append(sp)
    p.append(pPr)
    return p

def caption_el(md_text):
    m = re.match(r'^\*Hình (\d+\.\d+):\s*(.+?)\*$', md_text.strip())
    label, name = (m.group(1), m.group(2)) if m else (None, md_text.strip('*'))
    p = doc.add_paragraph()
    add_el(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    add_inline(p, f'Hình {label} – {name}', size=12)
    return p._p

def caption(md_line):
    m = re.match(r'^\*Hình (\d+\.\d+):\s*(.+?)\*$', md_line.strip())
    label, name = (m.group(1), m.group(2)) if m else (None, md_line.strip('*'))
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, indent=None, after=12, line=1.15)
    add_inline(p, f'Hình {label} – {name}', size=12, italic=False)
    return p

def shade_cell(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hexcolor}"/>'))

def set_tbl_props(tbl, font_size=12, line_sp=1.15, cant_split=True):
    tbl.style = doc.styles['Table Grid']
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:top w:w="40" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
        f'<w:bottom w:w="40" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>'))
    rows = tbl.rows
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        if cant_split:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for i, row in enumerate(rows):
        for cell in row.cells:
            if i == 0:
                shade_cell(cell, 'F2F2F2')
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = line_sp
                for r in p.runs:
                    r.font.size = Pt(font_size)
                    if i == 0:
                        r.font.bold = True

def add_md_table(rows):
    ncols = max(len(r) for r in rows)
    maxcell = max(len(c) for r in rows for c in r)
    header0 = rows[0][0].strip() if rows and rows[0] else ''
    is_phancong = (header0 == 'STT' and ncols == 5)
    fs = 10 if is_phancong else (11 if (maxcell > 400 or len(rows) >= 15) else 12)
    line_sp = 1.05 if is_phancong else 1.15
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    add_el(tbl._tbl)
    for i, r in enumerate(rows):
        for j in range(ncols):
            txt = r[j] if j < len(r) else ''
            cell = tbl.rows[i].cells[j]
            parts = txt.split('<br>')
            first = True
            for part in parts:
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                if part.strip():
                    add_inline(p, part.strip(), size=fs)
    set_tbl_props(tbl, fs, line_sp, cant_split=not is_phancong)
    if is_phancong:
        # cố định độ rộng cột: STT, Họ tên, Mã SV, Nhiệm vụ, Tỷ lệ (tổng 8504 twips = 15cm)
        widths = [600, 1700, 1100, 4304, 800]
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        for el in tblPr.findall(qn('w:tblW')):
            tblPr.remove(el)
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="8504" w:type="dxa"/>'))
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        grid = tbl._tbl.find(qn('w:tblGrid'))
        for gc in list(grid.findall(qn('w:gridCol'))):
            grid.remove(gc)
        for w in widths:
            grid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{w}"/>'))
        for row in tbl.rows:
            for j, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                for el in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(el)
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{widths[j]}" w:type="dxa"/>'))
    return tbl

def ascii_block(lines):
    maxlen = max(len(l) for l in lines) if lines else 0
    fs = 8.5 if maxlen <= 80 else (8 if maxlen <= 90 else 7)
    for ln in lines:
        p = doc.add_paragraph()
        add_el(p._p)
        pf = p.paragraph_format
        pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F7F7F7"/>'))
        r = p.add_run(ln if ln else ' ')
        set_run(r, fs, mono=True)

last_heading_el = [None]
def heading(text, level, page_break=False):
    p = doc.add_paragraph(style=f'Heading {level}')
    add_el(p._p)
    add_inline(p, text, size={1: 16, 2: 14, 3: 13}[level], bold=True)
    if page_break:
        p.paragraph_format.page_break_before = True
    last_heading_el[0] = p._p
    return p

def fm_title(text, page_break=True):
    return para(text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                indent=None, before=6, after=14, page_break=page_break)

# ---------- rebuild cover ----------
for el in cover_els:
    add_el(copy.deepcopy(el))

# thu nhỏ logo bìa để bìa gọn trong 1 trang
SCALE = 0.65
for el in body.iter(qn('wp:extent')):
    el.set('cx', str(int(int(el.get('cx')) * SCALE)))
    el.set('cy', str(int(int(el.get('cy')) * SCALE)))
for el in body.iter(qn('a:ext')):
    if el.get('cx') and int(el.get('cx')) > 1000000:  # chỉ ảnh lớn (logo)
        el.set('cx', str(int(int(el.get('cx')) * SCALE)))
        el.set('cy', str(int(int(el.get('cy')) * SCALE)))

# locate & rewrite title paragraph inside copied cover
for p in doc.paragraphs:
    if 'XÂY DỰNG WEBSITE' in p.text:
        # clear all runs
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        r = p.add_run('ĐỀ TÀI:'); set_run(r, 14, True)
        # add two more centered bold lines after this paragraph
        base = p._p
        for line in ('XÂY DỰNG WEBSITE BÁN LINH KIỆN ĐIỆN TỬ',
                     'TÍCH HỢP TRỢ LÝ KỸ THUẬT AI (ECSHOP)'):
            newp = copy.deepcopy(p._p)
            for rr in newp.findall(qn('w:r')):
                newp.remove(rr)
            base.addnext(newp)
            from docx.text.paragraph import Paragraph
            np = Paragraph(newp, p._parent)
            rr = np.add_run(line); set_run(rr, 16, True)
            base = newp
        break

# add khoa/khoa-hoc + place/date lines after the GV paragraph
for p in doc.paragraphs:
    if 'Giảng viên hướng dẫn' in p.text:
        base = p._p
        extra = ['Khoa: Công nghệ Thông tin – Trường Công nghệ Thông tin',
                 'Khóa học: K17 (2023 – 2027)']
        for line in extra:
            newp = copy.deepcopy(p._p)
            for rr in newp.findall(qn('w:r')):
                newp.remove(rr)
            base.addnext(newp)
            from docx.text.paragraph import Paragraph
            np = Paragraph(newp, p._parent)
            rr = np.add_run(line); set_run(rr, 14, True)
            base = newp
        # Hà Nội – 2026 centered, sát dưới khối GV/Khoa
        last = base
        e = OxmlElement('w:p')
        np = Paragraph(e, p._parent)
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np.paragraph_format.space_before = Pt(6)
        rr = np.add_run('Hà Nội – 2026'); set_run(rr, 14, True)
        last.addnext(e)
        break

# ---------- MD preprocessing: mermaid -> placeholders ----------
names = ['erd', 'usecase', 'act_order', 'act_jwt', 'act_ai',
         'seq_payment', 'seq_tools', 'cls_auth', 'cls_product', 'cls_ai']
it = iter(names)
MD = re.sub(r'```mermaid\n.*?```', lambda m: '@@IMG:' + next(it) + '@@', MD, flags=re.S)

lines = MD.split('\n')

# ---------- stream content ----------
CAPRE = re.compile(r'^\*Hình \d+\.\d+:.*\*$')
i = 0
n = len(lines)
started = False
skip_section = False
while i < n:
    raw = lines[i]
    line = raw.rstrip()
    s = line.strip()

    if not s or s == '---':
        i += 1; continue
    if not started:
        if re.match(r'^#{1,3}\s+LỜI CẢM ƠN\s*$', s):
            started = True
        else:
            i += 1; continue
    # skip the markdown pre-title block until LỜI CẢM ƠN
    if s.startswith('#'):
        m = re.match(r'^(#{1,3})\s+(.*)$', s)
        hashes, htext = m.group(1), m.group(2).strip()
        if htext == 'MỤC LỤC TỔNG QUAN':
            skip_section = True
            i += 1; continue
        if skip_section:
            skip_section = False  # this heading ends the skipped block
        if htext == 'LỜI CẢM ƠN':
            fm_title('LỜI CẢM ƠN')
        elif htext == 'LỜI CAM ĐOAN':
            fm_title('LỜI CAM ĐOAN')
        elif htext.startswith('BẢNG PHÂN CÔNG'):
            fm_title('BẢNG PHÂN CÔNG NHIỆM VỤ VÀ TỶ LỆ ĐÓNG GÓP')
        elif htext.startswith('DANH MỤC TỪ VIẾT TẮT'):
            fm_title('DANH MỤC TỪ VIẾT TẮT VÀ THUẬT NGỮ CHUYÊN NGÀNH')
        elif htext.startswith('DANH MỤC BẢNG BIỂU'):
            fm_title('DANH MỤC BẢNG BIỂU VÀ BIỂU ĐỒ MINH HỌA')
        elif htext == 'LỜI MỞ ĐẦU':
            heading('LỜI MỞ ĐẦU', 1, page_break=True)
        elif htext.startswith('TÀI LIỆU THAM KHẢO'):
            heading('TÀI LIỆU THAM KHẢO', 1, page_break=True)
        else:
            level = len(hashes)
            lvl = 1 if level == 1 else (2 if level == 2 else 3)
            heading(htext, lvl, page_break=(level == 1))
        i += 1; continue

    if skip_section:
        i += 1; continue

    # image placeholder
    m = re.match(r'^@@IMG:(\w+)@@$', s)
    if m:
        if m.group(1) in ('seq_payment', 'seq_tools'):
            # trang ngang: break đặt TRƯỚC heading 4.4/4.5 để heading + hình + caption cùng trang ngang
            pa = section_break_para(landscape=False, restart=(m.group(1) == 'seq_payment'))
            if last_heading_el[0] is not None:
                last_heading_el[0].addprevious(pa)
            else:
                add_el(pa)
            add_figure(m.group(1), landscape=True)
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            if j < n and CAPRE.match(lines[j].strip()):
                add_el(caption_el(lines[j].strip()))
                i = j
            add_el(section_break_para(landscape=True))
        else:
            add_figure(m.group(1))
        i += 1; continue

    # figure caption
    if re.match(r'^\*Hình \d+\.\d+:.*\*$', s):
        caption(s)
        i += 1; continue

    # ASCII block
    if s.startswith('```'):
        block = []
        i += 1
        while i < n and not lines[i].strip().startswith('```'):
            block.append(lines[i].rstrip('\n'))
            i += 1
        i += 1
        ascii_block(block)
        continue

    # table
    if s.startswith('|'):
        rows = []
        while i < n and lines[i].strip().startswith('|'):
            row_line = lines[i].strip()
            if not re.match(r'^\|[\s:\-|]+\|$', row_line):
                cells = [c.strip() for c in row_line.strip('|').split('|')]
                rows.append(cells)
            i += 1
        if rows:
            add_md_table(rows)
            para('', indent=None, after=6)
        continue

    # cam đoan signature block
    if s == '*Hà Nội, ngày 27 tháng 10 năm 2026*':
        para('Hà Nội, ngày 27 tháng 10 năm 2026', align=WD_ALIGN_PARAGRAPH.RIGHT,
             indent=None, before=18, after=2)
        i += 1; continue
    if s == '**Tập thể sinh viên thực hiện**':
        para('Tập thể sinh viên thực hiện', align=WD_ALIGN_PARAGRAPH.RIGHT,
             indent=None, after=2)
        i += 1; continue
    if s == '*Đàm Anh Pháp – Phạm Đức Tài*':
        para('Đàm Anh Pháp – Phạm Đức Tài', italic=True,
             align=WD_ALIGN_PARAGRAPH.RIGHT, indent=None, after=2)
        i += 1; continue

    # list items
    m = re.match(r'^(\s*)-\s+(.*)$', raw)
    if m:
        depth = len(m.group(1)) // 2
        p = para(indent=Cm(0.75 + 0.6 * depth), hanging=True, after=4)
        add_inline(p, '–  ' + m.group(2))
        i += 1; continue
    m = re.match(r'^(\d+)\.\s+(.*)$', s)
    if m:
        p = para(indent=Cm(0.75), hanging=True, after=4)
        add_inline(p, f'{m.group(1)}. {m.group(2)}')
        i += 1; continue

    # normal paragraph
    para(s)
    i += 1

# ---------- insert phiếu chấm thi (after phân công table) + section break after it ----------
target = None
for tbl in doc.tables:
    if 'TỔNG CỘNG MỨC ĐỘ ĐÓNG GÓP' in tbl._tbl.xml:
        target = tbl._tbl
        break
if target is not None:
    anchor = target
    last_phieu = None
    for el in phieu_els:
        cp = copy.deepcopy(el)
        anchor.addnext(cp)
        anchor = cp
        last_phieu = cp
    # section break right after the phiếu block (starts MỤC LỤC on a new page, section 2)
    cp = copy.deepcopy(sect_break_el)
    last_phieu.addnext(cp)

# ---------- MỤC LỤC TOC (before DANH MỤC TỪ VIẾT TẮT title) ----------
W = qn
def make_para():
    return OxmlElement('w:p')

title_p = make_para()
pPr = OxmlElement('w:pPr')
pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="240"/>'))
title_p.append(pPr)
r = OxmlElement('w:r')
rPr = OxmlElement('w:rPr')
rPr.append(OxmlElement('w:b'))
rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="28"/>'))
r.append(rPr)
t = OxmlElement('w:t'); t.text = 'MỤC LỤC'
r.append(t)
title_p.append(r)

toc_p = make_para()
def fld_run(child_builder):
    r = OxmlElement('w:r')
    r.append(child_builder())
    return r
fb = OxmlElement('w:fldChar'); fb.set(qn('w:fldCharType'), 'begin'); fb.set(qn('w:dirty'), 'true')
toc_p.append(fld_run(lambda: fb))
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
instr.text = ' TOC \\o "1-3" \\h \\z \\u '
toc_p.append(fld_run(lambda: instr))
fs_ = OxmlElement('w:fldChar'); fs_.set(qn('w:fldCharType'), 'separate')
toc_p.append(fld_run(lambda: fs_))
rnote = OxmlElement('w:r')
rPr2 = OxmlElement('w:rPr')
rPr2.append(OxmlElement('w:i'))
rPr2.append(parse_xml(f'<w:color {nsdecls("w")} w:val="808080"/>'))
rnote.append(rPr2)
tn = OxmlElement('w:t'); tn.text = '(Nhấp chuột phải vào mục lục và chọn "Update Field" để cập nhật số trang)'
rnote.append(tn)
toc_p.append(rnote)
fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end')
toc_p.append(fld_run(lambda: fe))

for p in doc.paragraphs:
    if p.text.strip() == 'DANH MỤC TỪ VIẾT TẮT VÀ THUẬT NGỮ CHUYÊN NGÀNH':
        p._p.addprevious(title_p)
        p._p.addprevious(toc_p)
        break

# numbering: chỉ restart ở section thân bài, phần cuối chạy liên tục
for pgnt in last_sectpr.findall(qn('w:pgNumType')):
    last_sectpr.remove(pgnt)

# ---------- ensure both sections show the page-number header ----------
hdr = None
for s_ in doc.sections:
    for hr in s_._sectPr.findall(qn('w:headerReference')):
        hdr = hr
        break
    if hdr is not None:
        break
for s_ in doc.sections:
    if s_._sectPr.find(qn('w:headerReference')) is None and hdr is not None:
        s_._sectPr.insert(0, copy.deepcopy(hdr))

doc.save(OUT)
print('saved', OUT)
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
